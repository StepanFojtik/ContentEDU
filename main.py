# project_structure/
# ├── documents/
# │   └── methodology.txt (tvůj metodologický soubor)
# ├── main.py
# └── requirements.txt

# ------------------ main.py -------------------
import os
# --- Nastavení API klíče ---
os.environ["OPENAI_API_KEY"] = ""  # <--- ZDE vlož platný klíč
import gradio as gr
from langchain_community.document_loaders import PyMuPDFLoader
from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain_openai import OpenAIEmbeddings, ChatOpenAI
from langchain_community.vectorstores import Chroma
from docx import Document

VECTOR_DB_PATH = "vectorstore"
conversation_history = []

def load_documents(paths):
    docs = []
    for path in paths:
        loader = PyMuPDFLoader(path)
        docs.extend(loader.load())
    return docs

def create_vector_db(docs):
    splitter = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=200)
    splits = splitter.split_documents(docs)
    embeddings = OpenAIEmbeddings()
    vectorstore = Chroma.from_documents(splits, embeddings, persist_directory=VECTOR_DB_PATH)
    vectorstore.persist()

def get_retriever():
    embeddings = OpenAIEmbeddings()
    vectorstore = Chroma(persist_directory=VECTOR_DB_PATH, embedding_function=embeddings)
    return vectorstore.as_retriever(search_type="mmr", search_kwargs={"k": 12})

def generate_course_outline(course_title, methodology, additional_info):
    llm = ChatOpenAI(model="gpt-3.5-turbo")
    prompt = f"""
    Vytvoř osnovu kurzu s názvem "{course_title}" na základě této metodiky:

    {methodology}

    Doplňující informace: {additional_info or "Bez dalších požadavků."}

    Osnova by měla obsahovat konkrétní názvy modulů a kapitol (žádné "Modul 1"), uveď hierarchii a strukturu výuky.
    """
    return llm.predict(prompt)

def generate_lesson_contents(outline, retriever):
    llm = ChatOpenAI(model="gpt-3.5-turbo")
    full_course = ""
    full_context = ""

    for line in outline.split("\n"):
        stripped = line.strip()
        if not stripped or stripped.lower().startswith("název"):
            full_course += line + "\n"
            continue

        relevant_docs = retriever.get_relevant_documents(stripped)
        context = "\n".join(doc.page_content for doc in relevant_docs)
        full_context += f"\n--- {stripped} ---\n{context}\n"

        prompt = f"""
        Téma kapitoly: {stripped}

        Na základě následujícího kontextu z dokumentů:

        {context}

        Vytvoř ucelený obsah této kapitoly pro univerzitní kurz. Zaměř se na:
        - teoretický výklad
        - definice
        - příklady
        - praktické aplikace

        Použij výukový styl, který je vhodný pro studium.
        """
        content = llm.predict(prompt)
        full_course += f"{line}\n{content}\n\n"

    return full_course, full_context

def export_to_word(content, filename="generated_course.docx"):
    doc = Document()
    doc.add_heading("Vygenerovaný kurz", 0)
    doc.add_paragraph(content)
    doc.save(filename)
    return filename

def export_to_html(content, filename="generated_course.html"):
    html_content = f"""<!DOCTYPE html>
<html lang='cs'>
<head>
    <meta charset='UTF-8'>
    <title>Vygenerovaný kurz</title>
</head>
<body>
    <h1>Vygenerovaný kurz</h1>
    <pre>{content}</pre>
</body>
</html>"""
    with open(filename, "w", encoding="utf-8") as f:
        f.write(html_content)
    return filename

def process(pdf_files, syllabus_file, course_title, additional_info):
    all_files = [f.name for f in pdf_files] + [syllabus_file.name, "documents/methodology.txt"]
    docs = load_documents(all_files)
    create_vector_db(docs)
    retriever = get_retriever()

    with open("documents/methodology.txt", "r", encoding="utf-8") as file:
        methodology = file.read()

    course_outline = generate_course_outline(course_title, methodology, additional_info)
    full_course, full_context = generate_lesson_contents(course_outline, retriever)

    conversation_history.append({"dotaz": course_outline, "odpoved": full_course})
    return full_course, full_context, conversation_history

def refine_course(feedback):
    if not conversation_history:
        return "Nejdříve vytvoř kurz.", conversation_history

    last_response = conversation_history[-1]['odpoved']
    llm = ChatOpenAI(model="gpt-3.5-turbo")
    new_prompt = f"Uprav následující obsah kurzu na základě připomínek:\n{feedback}\n\nPůvodní obsah:\n{last_response}"
    refined_content = llm.predict(new_prompt)

    conversation_history.append({"dotaz": feedback, "odpoved": refined_content})
    return refined_content, conversation_history

# --- Gradio UI ---
with gr.Blocks() as iface:
    gr.Markdown("# 🧠 Generátor Moodle kurzů")

    with gr.Row():
        pdf_files = gr.File(label="PDF dokumenty", file_count="multiple")
        syllabus_file = gr.File(label="Sylabus (PDF)")

    course_title = gr.Textbox(label="Název kurzu")
    additional_info = gr.Textbox(label="Doplňující info (nepovinné)")

    generate_btn = gr.Button("Vygenerovat kurz")
    course_output = gr.TextArea(label="📘 Vygenerovaný obsah kurzu")
    context_output = gr.TextArea(label="🔍 Použitý kontext z dokumentů")

    export_btn = gr.Button("Export do Wordu")
    export_output = gr.File(label="Stáhnout Word (DOCX)")

    export_html_btn = gr.Button("Export do HTML")
    export_html_output = gr.File(label="Stáhnout HTML")

    feedback = gr.Textbox(label="Připomínky")
    refine_btn = gr.Button("Upravit podle připomínek")
    conversation_display = gr.JSON(label="Historie")

    generate_btn.click(process, [pdf_files, syllabus_file, course_title, additional_info],
                       [course_output, context_output, conversation_display])
    export_btn.click(lambda content: export_to_word(content), inputs=course_output, outputs=export_output)
    export_html_btn.click(lambda content: export_to_html(content), inputs=course_output, outputs=export_html_output)
    refine_btn.click(refine_course, inputs=feedback, outputs=[course_output, conversation_display])

iface.launch()
